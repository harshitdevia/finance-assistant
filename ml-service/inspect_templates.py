import sys
import docx

def inspect_file(filename):
    print(f"=== INSPECTING {filename} ===")
    doc = docx.Document(filename)
    
    print("\n--- Paragraphs ---")
    for i, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        if text:
            # Print paragraph index and text if it's a heading or looks like a template guide
            if len(text) < 80 or any(kw in text for kw in ['Chapter', 'Section', 'Abstract', 'Executive Summary', 'Introduction', 'Review', 'Methodology', 'Conclusion', 'Bibliography', 'USN', 'Student Name']):
                print(f"Paragraph {i}: {text}")
            
    print("\n--- Tables ---")
    for t_idx, table in enumerate(doc.tables):
        print(f"Table {t_idx}:")
        for r_idx, row in enumerate(table.rows[:20]):
            row_texts = [cell.text.strip() for cell in row.cells]
            print(f"  Row {r_idx}: {row_texts}")

if __name__ == '__main__':
    if len(sys.argv) < 2:
        print("Usage: python inspect_templates.py <docx_file>")
        sys.exit(1)
    inspect_file(sys.argv[1])
